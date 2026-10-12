import argparse
import json
import sqlite3
import time
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

import config


FILTERS = {"all", "direct", "co-exhibitors"}
PARTICIPANT_ENDPOINT = "api/v1.0/participant/getAllParticipantList"
DETAIL_ENDPOINT = "api/v1.0/participant/getCompanyDetailBySlug"


def request_json(session, endpoint, payload):
    url = config.API_BASE_URL + endpoint
    for attempt in range(config.REQUEST_RETRIES + 1):
        response = session.post(
            url,
            json=payload,
            headers={
                "Accept": "application/json",
                "Content-Type": "application/json",
                "User-Agent": "ankiros-sync/1.0",
            },
            timeout=config.REQUEST_TIMEOUT,
        )
        if response.status_code != 429:
            response.raise_for_status()
            body = response.json()
            if body.get("hasError"):
                raise RuntimeError(
                    body.get("message") or f"ANKIROS API error: {endpoint}"
                )
            return body.get("result") or []
        if attempt == config.REQUEST_RETRIES:
            response.raise_for_status()
        retry_after = response.headers.get("Retry-After")
        try:
            delay = max(float(retry_after), config.REQUEST_DELAY)
        except (TypeError, ValueError):
            delay = config.REQUEST_DELAY * (2 ** attempt)
        time.sleep(delay)
    raise RuntimeError("ANKIROS request retry loop completed unexpectedly.")


def fetch_exhibitors(session=None):
    """Fetch direct exhibitors and expand their co-exhibitor details."""

    session = session or requests.Session()
    direct = request_json(
        session,
        PARTICIPANT_ENDPOINT,
        {"expoId": config.EXPO_ID},
    )
    records = []
    for item in direct:
        parent = {
            "company_name": item.get("companyName", "").strip(),
            "slug": item.get("slug"),
            "relationship": "direct",
            "parent_company": None,
        }
        records.append(parent)
        if not item.get("slug"):
            continue
        detail = request_json(
            session,
            DETAIL_ENDPOINT,
            {"expoId": config.EXPO_ID, "slug": item["slug"]},
        )
        for co_exhibitor in detail.get("subcoCompany", []) or []:
            if not isinstance(co_exhibitor, dict):
                continue
            records.append(
                {
                    "company_name": (
                        co_exhibitor.get("companyName")
                        or co_exhibitor.get("name")
                        or ""
                    ).strip(),
                    "slug": co_exhibitor.get("slug"),
                    "relationship": "co-exhibitor",
                    "parent_company": parent["company_name"],
                }
            )
    return records


def filter_exhibitors(records, relationship_filter="all"):
    """Filter direct and co-exhibitor records without changing their data."""

    if relationship_filter not in FILTERS:
        raise ValueError(
            f"Unknown filter {relationship_filter!r}; choose from "
            f"{', '.join(sorted(FILTERS))}"
        )
    if relationship_filter == "all":
        return list(records)
    relationship = (
        "direct" if relationship_filter == "direct" else "co-exhibitor"
    )
    return [
        record for record in records
        if record.get("relationship") == relationship
    ]


def initialize_database(path=config.DATABASE_PATH):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(path) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS exhibitors (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_name TEXT NOT NULL,
                slug TEXT,
                relationship TEXT NOT NULL,
                parent_company TEXT,
                edition INTEGER NOT NULL,
                first_seen_at TEXT NOT NULL,
                UNIQUE(company_name, relationship, parent_company, edition)
            )
            """
        )


def sync_records(records, path=config.DATABASE_PATH):
    initialize_database(path)
    now = datetime.now(timezone.utc).isoformat()
    with sqlite3.connect(path) as connection:
        for record in records:
            connection.execute(
                """
                INSERT INTO exhibitors (
                    company_name, slug, relationship, parent_company,
                    edition, first_seen_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(company_name, relationship, parent_company, edition)
                DO UPDATE SET slug = excluded.slug
                """,
                (
                    record["company_name"],
                    record.get("slug"),
                    record["relationship"],
                    record.get("parent_company"),
                    config.EDITION,
                    now,
                ),
            )


def export_records(records, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Exhibitors"
    worksheet.append([
        "Company Name", "Relationship", "Parent Company", "Slug", "Edition"
    ])
    for record in records:
        worksheet.append([
            record["company_name"],
            record["relationship"],
            record.get("parent_company"),
            record.get("slug"),
            config.EDITION,
        ])
    worksheet.freeze_panes = "A2"
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    if worksheet.max_row > 1:
        table = Table(displayName="AnkirosExhibitors", ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    for column in worksheet.columns:
        worksheet.column_dimensions[column[0].column_letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2,
            50,
        )
    workbook.save(output_path)
    return output_path


def main(argv=None):
    parser = argparse.ArgumentParser(description="Sync ANKIROS exhibitors")
    subparsers = parser.add_subparsers(dest="command", required=True)
    for command in ("sync", "export"):
        subparser = subparsers.add_parser(command)
        subparser.add_argument("--filter", choices=sorted(FILTERS), default="all")
    args = parser.parse_args(argv)
    records = filter_exhibitors(
        fetch_exhibitors(),
        args.filter,
    )
    if args.command == "sync":
        sync_records(records)
    output = config.EXPORTS_DIR / f"ANKIROS_{config.EDITION}.xlsx"
    export_records(records, output)
    print(f"{len(records)} {args.filter} exhibitors processed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
