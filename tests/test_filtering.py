import sqlite3

import pytest
from openpyxl import load_workbook

from ankiros_sync import (
    export_records,
    fetch_exhibitors,
    filter_exhibitors,
    sync_records,
)


RECORDS = [
    {
        "company_name": "Direct A",
        "country": "Germany",
        "relationship": "direct",
        "parent_company": None,
    },
    {
        "company_name": "Co A",
        "country": "Türkiye",
        "relationship": "co-exhibitor",
        "parent_company": "Direct A",
    },
]


def test_filter_returns_direct_exhibitors_only():
    assert [r["company_name"] for r in filter_exhibitors(RECORDS, "direct")] == [
        "Direct A"
    ]


def test_filter_returns_co_exhibitors_only():
    assert [
        r["company_name"]
        for r in filter_exhibitors(RECORDS, "co-exhibitors")
    ] == ["Co A"]


def test_filter_all_preserves_both_relationships():
    assert filter_exhibitors(RECORDS, "all") == RECORDS


def test_filter_rejects_unknown_value():
    with pytest.raises(ValueError):
        filter_exhibitors(RECORDS, "partners")


def test_fetch_expands_subco_companies():
    class Session:
        def post(self, url, **kwargs):
            class Response:
                status_code = 200
                headers = {}

                def raise_for_status(self):
                    return None

                def json(self):
                    if url.endswith("getAllParticipantList"):
                        return {
                            "result": [
                                {
                                    "companyName": "Direct A",
                                    "countryName": "Germany",
                                    "slug": "direct-a",
                                }
                            ]
                        }
                    if kwargs["json"]["slug"] == "direct-a":
                        return {
                            "result": {
                                "address": "Street 1",
                                "country": "Germany",
                                "standNo": "A01",
                                "phone": "+49 123",
                                "email": "direct@example.com",
                                "exhibitList": [
                                    {"titleEng": "Casting"},
                                ],
                                "subcoCompany": [
                                    {
                                        "subcoCompanyName": "Co A",
                                        "country": "Türkiye",
                                        "slug": "co-a",
                                    }
                                ],
                            }
                        }
                    return {
                        "result": {
                            "address": "Street 2",
                            "standNo": "A02",
                            "phone": "+49 456",
                            "email": "co@example.com",
                            "exhibitList": [{"titleEng": "Furnaces"}],
                            "subcoCompany": [],
                        }
                    }

            return Response()

    records = fetch_exhibitors(Session())
    assert records == [
        {
            "company_name": "Direct A",
            "country": "Germany",
            "slug": "direct-a",
            "relationship": "direct",
            "parent_company": None,
            "address": "Street 1",
            "stand_no": "A01",
            "phone": "+49 123",
            "email": "direct@example.com",
            "products": "Casting",
        },
        {
            "company_name": "Co A",
            "country": "Türkiye",
            "slug": "co-a",
            "relationship": "co-exhibitor",
            "parent_company": "Direct A",
            "address": "Street 2",
            "stand_no": "A02",
            "phone": "+49 456",
            "email": "co@example.com",
            "products": "Furnaces",
        },
    ]


def test_sync_persists_detail_fields(tmp_path):
    records = [{
        "company_name": "Direct A",
        "country": "Germany",
        "slug": "direct-a",
        "relationship": "direct",
        "parent_company": None,
        "address": "Street 1",
        "stand_no": "A01",
        "phone": "+49 123",
        "email": "direct@example.com",
        "products": "Casting",
    }]
    database_path = tmp_path / "ankiros.db"

    sync_records(records, database_path)

    with sqlite3.connect(database_path) as connection:
        row = connection.execute(
            "SELECT country, address, stand_no, phone, email, products "
            "FROM exhibitors"
        ).fetchone()
    assert row == (
        "Germany",
        "Street 1",
        "A01",
        "+49 123",
        "direct@example.com",
        "Casting",
    )


def test_export_matches_sales_leads_shape(tmp_path):
    output = export_records(
        [{
            "company_name": "Direct A",
            "country": "Germany",
            "relationship": "direct",
            "parent_company": None,
            "slug": "direct-a",
            "address": "Street 1\x1e",
            "stand_no": "A01",
            "phone": "+49 123",
            "email": "direct@example.com",
            "products": "Casting",
        }],
        tmp_path / "ANKIROS_2026.xlsx",
    )

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["Summary", "Company Master", "Sales Leads"]
    assert list(workbook["Sales Leads"].iter_rows(max_row=1, values_only=True))[0] == (
        "Company Name",
        "Relationship",
        "Parent Company",
        "Address",
        "Country",
        "Stand",
        "Phone",
        "Email",
        "Product Categories",
        "Slug",
        "Edition",
    )
    assert workbook["Sales Leads"].cell(row=2, column=5).value == "Germany"
